#!/usr/bin/env python3
"""
Turns a lead-analysis workbook (ClickTap_Lead_Analysis.xlsx) into one SQL file
you paste into Supabase -> SQL Editor. See docs/IMPORT.md.

    python3 scripts/import_lead_analysis.py ClickTap_Lead_Analysis.xlsx imports/leads.sql

What it creates (all tagged "import-<batch>" so it can be undone in one line):
  * one company per business (repeat inquiries merged by email domain / phone)
  * its contacts
  * HOT/WARM leads: an open deal at the right stage with a next follow-up date,
    and the pitch / opening message / next action in the deal notes
  * existing clients: a "won" deal for what they bought + an open upsell deal
  * every original inquiry as a timeline activity, with your call notes
  * COLD leads: company + contact only, tagged "cold" (no deal)
  * "Not interested / invalid": imported as archived (kept, but hidden)

The output contains personal data. Never commit it to git.
Needs: pip install openpyxl
"""
import datetime
import re
import sys
import unicodedata
import uuid

import openpyxl

BATCH = "import-2026-10"
FIRST_FOLLOW_UP = datetime.date(2026, 10, 27)  # HOT leads: the next working day
PER_DAY = 11  # other WARM leads: how many follow-ups to schedule per working day


def working_days(start):
    """Mon-Fri from `start` onwards (UAE weekend is Saturday-Sunday)."""
    day = start
    while True:
        if day.weekday() < 5:
            yield day
        day += datetime.timedelta(days=1)

# Short service names -> (deal title prefix, services-catalogue name, billing)
SERVICES = {
    "PR & media publications": ("PR", "PR & media publications", "one_off"),
    "SEO & organic Google ranking": ("SEO", "SEO", "monthly"),
    "Google Ads & lead-generation campaigns": ("Google Ads", "Google ads management", "monthly"),
    "Google Ads & lead-generation campaigns (Search + Shopping)": ("Google Ads", "Google ads management", "monthly"),
    "Website design & redevelopment": ("Website", "Website design", "one_off"),
    "Social media management & content creation": ("Social media", "Social media management", "monthly"),
    "Google Maps optimisation & review growth": ("Google Maps", "Google Maps optimisation", "monthly"),
}
JUNK = {"", "-", "--", ".", "ok", "nn", "na", "n/a", "no", "none", "nil", "self", "self-employed", "self employed",
        "personal", "individual", "test", "home", "x", "xx", "abc", "company", "own", "ow", "private", "freelancer",
        "nothing", "yes", "hi", "dubai", "uae", "business", "owner", "me", "mh", "sd", "hem", "vb", "xyz", "att",
        "plc", "events", "artist", "enterpreneur", "entrepreneur", "no company", "properties", "eco", "juno", "nova"}
FREEMAIL = re.compile(r"@(gmail|gnail|gamil|gmeil|gmai|hotmail|yahoo|outlook|icloud|live|msn|aol|ymail|proton|"
                      r"windowslive|mail|email|rediffmail|vahoo|my)\.", re.I)


def sql(v):
    """Python value -> SQL literal."""
    if v is None or v == "":
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, list):
        return "array[" + ",".join(sql(x) for x in v) + "]::text[]"
    return "'" + str(v).replace("'", "''") + "'"


def clean(s):
    return unicodedata.normalize("NFKC", str(s or "")).strip()


def phone_e164(raw):
    """Best effort: UAE numbers to +971..., others keep their country code."""
    first = re.split(r"[/,;]| or ", clean(raw))[0]
    d = re.sub(r"\D", "", first.split(".")[0])
    if d.startswith("00"):
        d = d[2:]
    if len(d) == 9 and d.startswith("5"):
        d = "971" + d
    if len(d) == 10 and d.startswith("05"):
        d = "971" + d[1:]
    return "+" + d if 10 <= len(d) <= 15 and not d.startswith("0") else None


def valid_email(e):
    e = clean(e).lower()
    return e if re.fullmatch(r"[^@\s]+@[^@\s]+\.[a-z]{2,}", e) else None


def usable_company(name):
    n = clean(name)
    return n if n.lower() not in JUNK and not n.startswith("Lead |") and len(n) > 2 and "?" not in n else None


def stage_for(notes, closed):
    n = notes.lower()
    if closed:
        return "won"
    if re.search(r"proposal sent|quotation sent|quotation shared|sent proposal|costings|proposal to be sent|"
                 r"sent pr packages|last rate given|got quotation", n):
        return "proposal"
    if re.search(r"meet(ing)? done|meeting done|had (a )?meeting|2nd meeting", n):
        return "discovery"
    if re.search(r"google meet|meeting fixed|meeting on|interested|rc|rate ?card|details|packages|call back|wants", n):
        return "contacted"
    return "lead"


def channel_for(notes):
    n = notes.lower()
    if re.search(r"google meet|meeting|gmeet", n):
        return "meeting"
    if "whatsapp" in n:
        return "whatsapp"
    if re.search(r"email|mailed", n):
        return "email"
    if n:
        return "call"
    return "other"


def main(src, out):
    wb = openpyxl.load_workbook(src, data_only=True)
    leads = [dict(zip([c.value for c in wb["All Leads Scored"][1]], r))
             for r in wb["All Leads Scored"].iter_rows(min_row=2, values_only=True) if r[0]]
    prio = {}
    for r in wb["Priority Leads"].iter_rows(min_row=2, values_only=True):
        row = dict(zip([c.value for c in wb["Priority Leads"][1]], r))
        for lid in str(row["Lead ID (your file)"]).split(","):
            prio[lid.strip()] = row

    # ---- Group inquiries into companies (by researched company, email domain or phone)
    groups, key_of = {}, {}
    for lead in sorted(leads, key=lambda x: str(x["Inquiry date"])):
        lid = lead["Lead ID"]
        keys = []
        if lid in prio:
            keys.append("prio:" + prio[lid]["Company"])
        if lead["Company email domain"]:
            keys.append("dom:" + lead["Company email domain"].lower())
        ph = phone_e164(lead["Phone"])
        if ph:
            keys.append("ph:" + ph)
        gid = next((key_of[k] for k in keys if k in key_of), None) or str(uuid.uuid4())
        for k in keys:
            key_of.setdefault(k, gid)
        groups.setdefault(gid, []).append(lead)

    temp_rank = {"HOT": 0, "WARM": 1, "COLD": 2}
    companies, contacts, deals, deal_services, activities = [], [], [], [], []
    warm_slot, rest_days, rest_slot = 0, list(), 0
    days = working_days(FIRST_FOLLOW_UP + datetime.timedelta(days=6))
    ordered = sorted(groups.items(), key=lambda g: min((temp_rank[l["Temperature"]], -(l["Score /100"] or 0)) for l in g[1]))
    for position, (gid, rows) in enumerate(ordered, start=1):
        best = min(rows, key=lambda l: (temp_rank[l["Temperature"]], -(l["Score /100"] or 0)))
        p = prio.get(best["Lead ID"])
        temp = best["Temperature"]
        category = best["Category"] or ""
        closed_row = next((l for l in rows if re.search(r"deal closed", l["Your call notes"] or "", re.I)), None)
        negative = all((l["Category"] or "") == "Not interested / invalid" for l in rows)
        partner = any((l["Category"] or "") == "Partner (agency/reseller)" for l in rows)

        # Company name
        name = (p and p["Company"]) or next((usable_company(l["Company (as entered)"]) for l in rows
                                             if usable_company(l["Company (as entered)"])), None)
        domain = next((l["Company email domain"] for l in rows if l["Company email domain"]), None)
        if not name and domain:
            name = domain.split(".")[0].replace("-", " ").title() + f" ({domain})"
        if not name:
            person = clean(best["Contact name"]) or "Unknown"
            name = f"{person} (individual)"

        tags = [BATCH, temp.lower(), f"{best['Source tab'].lower()} campaign"]
        if partner:
            tags.append("partner")
        if closed_row:
            tags.append("client")
        if any(l["Count once (1 = first inquiry)"] == 0 for l in rows) or len(rows) > 1:
            tags.append("repeat inquiry")

        notes = []
        if p:
            notes += [f"What they do: {p['What they do (researched)']}",
                      f"Research: {p['Online research findings']}",
                      f"Decision-maker: {p['Decision-maker (verified vs unverified)']}",
                      f"Best person: {p['Best person to approach & why']}",
                      f"Sources: {(p['Sources'] or '').replace(chr(10), ' ')}"]
        notes.append(f"Lead rating {temp} {best['Score /100']}/100. {best['Why this rating'] or ''}")
        flags = sorted({f for l in rows for f in (l["Data-quality flags"] or "").split("; ") if f})
        if flags:
            notes.append("Data flags: " + "; ".join(flags))
        notes.append("Imported from lead IDs: " + ", ".join(l["Lead ID"] for l in rows))

        website = None
        if domain and not FREEMAIL.search("@" + domain):
            website = "https://" + domain
        company_id = str(uuid.uuid4())
        companies.append(dict(
            id=company_id, name=name[:200], industry=(p["What they do (researched)"][:100] if p else None),
            website=website, source="other", tags=tags,
            status="archived" if negative else "client" if closed_row else "lead",
            notes="\n\n".join(notes), position=position))

        # Contacts (one per distinct phone/email)
        seen, contact_ids = {}, {}
        for l in rows:
            ph, em = phone_e164(l["Phone"]), valid_email(l["Email"])
            key = ph or em or clean(l["Contact name"])
            if key in seen:
                contact_ids[l["Lead ID"]] = seen[key]
                continue
            cname = clean(l["Contact name"])
            if not cname or cname.lower() in {"ceo", "na", "ok"} or cname.startswith("Lead |") or cname == name:
                cname = (em.split("@")[0].replace(".", " ").title() if em else "Contact") + f" ({clean(l['Job title']) or 'role unknown'})"
            cid = str(uuid.uuid4())
            seen[key] = contact_ids[l["Lead ID"]] = cid
            extra = []
            if l["Phone"] and not ph:
                extra.append(f"Phone as entered: {l['Phone']}")
            if l["Email"] and not em:
                extra.append(f"Email as entered: {l['Email']}")
            contacts.append(dict(id=cid, company_id=company_id, full_name=cname[:150],
                                 job_title=clean(l["Job title"])[:100] or None, phone=ph, email=em,
                                 notes="; ".join(extra) or None))
        best_contact = contact_ids.get(best["Lead ID"])

        # Deals
        deal_id = None
        if closed_row:
            svc = SERVICES["PR & media publications"]
            won_id = str(uuid.uuid4())
            deals.append(dict(id=won_id, company_id=company_id, primary_contact_id=contact_ids.get(closed_row["Lead ID"]),
                              title=f"PR – {name}"[:200], billing="one_off", stage="won", next_follow_up_date=None,
                              notes=f"Closed deal from the PR campaign. Call notes: {closed_row['Your call notes']}\nValue not recorded in the sheet – please update.",
                              position=position))
            deal_services.append((won_id, svc[1]))
        if temp in ("HOT", "WARM") and not partner:
            service = (p and p["Best service to pitch first"]) or best["Best service to pitch first"] or "PR & media publications"
            prefix, catalogue, billing = SERVICES.get(service, ("PR", "PR & media publications", "one_off"))
            if temp == "HOT":
                follow = FIRST_FOLLOW_UP
            elif p:
                follow = FIRST_FOLLOW_UP + datetime.timedelta(days=1 + (warm_slot % 3))
                warm_slot += 1
            else:
                # Spread the rest over the following weeks, PER_DAY per working day, best first.
                if rest_slot % PER_DAY == 0:
                    rest_days.append(next(days))
                follow = rest_days[-1]
                rest_slot += 1
            deal_id = str(uuid.uuid4())
            stage = "contacted" if closed_row else stage_for(" ".join(l["Your call notes"] or "" for l in rows), False)
            message = (p and p["Personalised opening message"]) or best["Opening message"] or ""
            next_action = (p and p["Next action"]) or best["Next action"] or ""
            upsell = (p and p["Upsell later"]) or best["Upsell later"] or ""
            why = (p and p["Why this is the strongest entry point"]) or ""
            problem = (p and p["Biggest marketing problem"]) or ""
            deal_notes = "\n\n".join(x for x in [
                f"Rating: {temp} {best['Score /100']}/100",
                f"Biggest problem: {problem}" if problem else "",
                f"Pitch first: {service}" + (f" – {why}" if why else ""),
                f"Upsell later: {upsell}" if upsell else "",
                f"Next action: {next_action}",
                f"Opening message:\n{message}",
            ] if x)
            deals.append(dict(id=deal_id, company_id=company_id, primary_contact_id=best_contact,
                              title=(f"{prefix} upsell – {name}" if closed_row else f"{prefix} – {name}")[:200],
                              billing=billing, stage=stage, next_follow_up_date=follow.isoformat(),
                              notes=deal_notes, position=position))
            deal_services.append((deal_id, catalogue))

        # Activities: one per original inquiry
        for l in rows:
            date = str(l["Inquiry date"] or "")[:10] or "2026-01-01"
            budget = f" Budget: {l['Budget']}." if l["Budget"] else ""
            summary = (f"Inbound lead from the {l['Source tab']} ad campaign ({l['Campaign'] or 'campaign name not recorded'}): "
                       f"asked about {l['Asked about']}.{budget}")
            call_notes = clean(l["Your call notes"])
            activities.append(dict(company_id=company_id, deal_id=deal_id, contact_id=contact_ids.get(l["Lead ID"]),
                                   occurred_at=f"{date} 12:00:00+04", channel=channel_for(call_notes),
                                   summary=summary + (f"\n\nCall notes: {call_notes}" if call_notes else "\n\nNo call notes recorded."),
                                   outcome=f"{l['Temperature']} {l['Score /100']}/100" + (f" – {l['Category']}" if l["Category"] else "")))

    write_sql(out, companies, contacts, deals, deal_services, activities)
    print(f"{out}: {len(companies)} companies, {len(contacts)} contacts, {len(deals)} deals, {len(activities)} activities")


def values(rows, cols):
    return ",\n".join("  (" + ", ".join(sql(r[c]) for c in cols) + ")" for r in rows)


def write_sql(out, companies, contacts, deals, deal_services, activities):
    owner = "(select id from public.profiles where role = 'owner' order by created_at limit 1)"
    parts = [f"""-- Lead import ({BATCH}). Paste into Supabase -> SQL Editor and click Run.
-- Undo everything this file created:
--   delete from public.companies where '{BATCH}' = any(tags);
-- Contains personal data: do not commit this file to git.

begin;

do $$
begin
  if not exists (select 1 from public.profiles where role = 'owner') then
    raise exception 'Create your owner login first (docs/SETUP.md step 5).';
  end if;
  if exists (select 1 from public.companies where '{BATCH}' = any(tags)) then
    raise exception 'These leads were already imported. Run the undo line at the top first if you want to re-import.';
  end if;
end $$;

-- Services used by the imported deals (added only if missing)
insert into public.services (name, billing, sort_order) values
  ('PR & media publications', 'one_off', 5),
  ('Google Maps optimisation', 'monthly', 55)
on conflict (name) do nothing;
"""]
    for i in range(0, len(companies), 200):
        chunk = companies[i:i + 200]
        parts.append("insert into public.companies (id, name, industry, website, source, tags, status, notes, owner_id, created_by)\n"
                     "select v.id::uuid, v.name, v.industry, v.website, v.source::public.company_source, v.tags,\n"
                     f"       v.status::public.company_status, v.notes, {owner}, {owner}\n"
                     "from (values\n" + values(chunk, ["id", "name", "industry", "website", "source", "tags", "status", "notes"]) +
                     "\n) as v(id, name, industry, website, source, tags, status, notes);\n")
    for i in range(0, len(contacts), 200):
        chunk = contacts[i:i + 200]
        parts.append("insert into public.contacts (id, company_id, full_name, job_title, phone, email, notes, created_by)\n"
                     f"select v.id::uuid, v.company_id::uuid, v.full_name, v.job_title, v.phone, v.email, v.notes, {owner}\n"
                     "from (values\n" + values(chunk, ["id", "company_id", "full_name", "job_title", "phone", "email", "notes"]) +
                     "\n) as v(id, company_id, full_name, job_title, phone, email, notes);\n")
    if deals:
        parts.append("insert into public.deals (id, company_id, primary_contact_id, title, billing, stage, next_follow_up_date, notes, position, owner_id, created_by)\n"
                     "select v.id::uuid, v.company_id::uuid, v.primary_contact_id::uuid, v.title, v.billing::public.billing_type,\n"
                     f"       v.stage::public.deal_stage, v.next_follow_up_date::date, v.notes, v.position * 1024, {owner}, {owner}\n"
                     "from (values\n" + values(deals, ["id", "company_id", "primary_contact_id", "title", "billing", "stage", "next_follow_up_date", "notes", "position"]) +
                     "\n) as v(id, company_id, primary_contact_id, title, billing, stage, next_follow_up_date, notes, position);\n")
        parts.append("insert into public.deal_services (deal_id, service_id)\n"
                     "select v.deal_id::uuid, s.id from (values\n" +
                     ",\n".join(f"  ({sql(d)}, {sql(s)})" for d, s in deal_services) +
                     "\n) as v(deal_id, service_name) join public.services s on s.name = v.service_name\non conflict do nothing;\n")
    for i in range(0, len(activities), 200):
        chunk = activities[i:i + 200]
        parts.append("insert into public.activities (company_id, deal_id, contact_id, occurred_at, channel, summary, outcome, created_by)\n"
                     "select v.company_id::uuid, v.deal_id::uuid, v.contact_id::uuid, v.occurred_at::timestamptz,\n"
                     f"       v.channel::public.activity_channel, v.summary, v.outcome, {owner}\n"
                     "from (values\n" + values(chunk, ["company_id", "deal_id", "contact_id", "occurred_at", "channel", "summary", "outcome"]) +
                     "\n) as v(company_id, deal_id, contact_id, occurred_at, channel, summary, outcome);\n")
    parts.append(f"""
-- Quick check: what was imported
select status, count(*) from public.companies where '{BATCH}' = any(tags) group by status;

commit;
""")
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
